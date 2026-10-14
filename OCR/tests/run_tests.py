"""Regression & acceptance tests for doc-rebuilder.html (headless Chromium, offline libs).
Usage: python3 tests/run_tests.py   → prints a table and writes tests/out/results.json"""
import sys, os, json, re, time, collections, subprocess
sys.path.insert(0, os.path.dirname(__file__))
from harness import App, OUT, FIX
import docx as pydocx, openpyxl, pdfplumber, jiwer

R = []
def check(name, ok, detail=''):
    R.append({'test': name, 'ok': bool(ok), 'detail': str(detail)[:300]}); print(('PASS ' if ok else 'FAIL ') + name + ('  — ' + str(detail)[:160] if detail else ''))
def norm(t):
    t = re.sub(r'[ً-ٰٟـ‎‏]', '', t)
    t = t.translate(str.maketrans('٠١٢٣٤٥٦٧٨٩أإآىة', '0123456789اااىه'))
    return ' '.join(re.findall(r'[\w.,()%/-]+', t))
def nums(t): return [x for x in re.findall(r'\(?\d[\d,]*\.?\d*\)?%?', t) if len(x.strip('(),.%')) >= 2]

truth = json.load(open(os.path.join(FIX, 'truth.json')))
a = App()
try:
    # ---------- text PDF, Arabic financial statements ----------
    a.open('fin_ar.pdf'); t0 = time.time(); a.run(); ps = a.pages()
    txt = '\n'.join(p['text'] for p in ps); ref_fin = [p['text'] for p in ps]
    missing = [n for n in truth['fin_ar']['numbers'] if n not in txt]
    check('textPDF-ar: all ground-truth amounts present', not missing, missing)
    check('textPDF-ar: negative amounts keep parentheses order', '(125,000)' in txt and ')125,000(' not in txt)
    check('textPDF-ar: lam-alef ligature repaired', 'الاستثمار' in txt and 'االستثمار' not in txt)
    check('textPDF-ar: mixed-direction English sentence intact', 'EGP 2,250,832 and the date 31/12/2025' in txt)
    h1 = a.page.evaluate("document.querySelector('.pg[data-n=\"1\"] .ed').innerHTML")
    check('textPDF-ar: ruled table with merged cells (colspan & rowspan)', 'colspan="2"' in h1 and 'rowspan="2"' in h1)
    check('textPDF-ar: header row shading recovered', re.search(r'background:#D', h1) is not None)
    check('textPDF-ar: two-column region detected on page 2', 'class="mc"' in ps[1]['html'])
    check('textPDF-ar: logo image extracted', '<img' in h1)
    check('textPDF-ar: hyperlink kept', 'href="https://fra.gov.eg' in h1)
    check('textPDF-ar: repeated header + page number roles', a.page.evaluate("document.querySelectorAll('[data-role=hf]').length") >= 2 and a.page.evaluate("document.querySelectorAll('[data-role=pn]').length") >= 2)
    rep = a.page.evaluate("(() => { const r = VALID.run(APP.cur); return { e: r.findings.filter(f => f.s === 'e').length, foot: r.foot }; })()")
    check('validation: no source/output digit mismatch on text PDF', rep['e'] == 0, rep)
    check('validation: footing check finds matching totals', rep['foot']['checked'] >= 3 and rep['foot']['matched'] == rep['foot']['checked'], rep['foot'])
    # edit a number → validation reports it
    a.page.evaluate("(() => { const ed = document.querySelector('.pg[data-n=\"1\"] .ed'); const w = document.createTreeWalker(ed, NodeFilter.SHOW_TEXT); let n; while ((n = w.nextNode())) if (n.data.includes('615,200')) { n.data = n.data.replace('615,200', '651,200'); break; } ed.dispatchEvent(new Event('input', { bubbles: true })); })()")
    msg = a.page.evaluate("VALID.run(APP.cur).findings.map(f => f.m).join('\\n')")
    check('validation: manual number change is reported', '615200' in msg and '651200' in msg, msg[:200])
    p = a.export('docx', 'preserve', 'T_fin_preserve.docx'); d = pydocx.Document(p)
    check('export DOCX(preserve): opens in python-docx, 2 tables, image, 2+ sections', len(d.tables) == 2 and len(d.inline_shapes) == 1 and len(d.sections) >= 2)
    xl = a.export('xlsx', fname='T_fin.xlsx'); wb = openpyxl.load_workbook(xl); ws = wb.worksheets[0]
    vals = [c for row in ws.iter_rows(values_only=True) for c in row]
    check('export XLSX: numbers numeric, parentheses → negative, RTL sheet', -125000 in vals and 9116525 in vals and ws.sheet_view.rightToLeft)
    xl2 = a.export('xlsx', fname='T_fin_check.xlsx'); wb2 = openpyxl.load_workbook(xl2)
    chk = list(wb2['تحقق المجاميع'].iter_rows(values_only=True))[1:]
    check('XLSX: footing sheet with live SUM formulas for every total (≥10)', len(chk) >= 10 and all(str(r[4]).startswith('=SUM(') for r in chk), len(chk))
    check('XLSX: source reference column per row', any('صفحة 1 · جدول 1' in str(c) for r in wb2.worksheets[0].iter_rows(values_only=True) for c in r))
    check('audit trail: number edit logged with before/after', a.page.evaluate("APP.cur.log.some(e => e.k === 'num' && /615200/.test(e.b) && /651200/.test(e.a))"))
    check('XLSX: audit-trail sheet included', 'سجل التعديلات' in wb2.sheetnames)
    ht = a.export('html', 'preserve', 'T_fin.html'); check('export HTML: standalone file with tables and embedded image', '<table' in open(ht, encoding='utf-8').read() and 'data:image/' in open(ht, encoding='utf-8').read())
    tx = a.export('txt', fname='T_fin.txt'); tt = open(tx, encoding='utf-8-sig').read(); check('export TXT: page breaks + amounts', '\f' in tt and '9,116,525' in tt)

    # ---------- invoice: multi-page table ----------
    a.open('invoice_en.pdf'); a.run(); ref_inv = [p['text'] for p in a.pages()]
    check('multi-page table: continuation flagged', a.page.evaluate("document.querySelectorAll('table[data-cont]').length") == 2)
    xl = a.export('xlsx', fname='T_inv.xlsx'); wb = openpyxl.load_workbook(xl)
    ws = max(wb.worksheets, key=lambda w: w.max_row)
    rows = [r for r in ws.iter_rows(values_only=True)]
    items = [r for r in rows if isinstance(r[0], (int, float))]
    check('multi-page table merged in XLSX: 55 item rows, single header', len(items) == 55 and sum(1 for r in rows if r[1] == 'Description') == 1, (len(items)))
    tot = [r for r in rows if r[0] == 'Total']
    check('XLSX: invoice total parsed as number', tot and abs(tot[0][5] - float(truth['invoice_en']['total'].replace(',', ''))) < 0.01, tot)
    amt_ok = sum(1 for r, t in zip(items, truth['invoice_en']['items']) if abs((r[5] or 0) - float(t[5].replace(',', ''))) < 0.005)
    check('XLSX: all 55 line amounts exact', amt_ok == 55, amt_ok)
    p = a.export('docx', 'preserve', 'T_inv.docx'); d = pydocx.Document(p)
    check('DOCX: repeated header rows marked (w:tblHeader)', 'tblHeader' in d.element.xml)

    # ---------- native DOCX ----------
    a.open('fin_ar.docx'); ps = a.pages(); allh = ''.join(p['html'] for p in ps)
    check('DOCX in: no OCR, merged cells, columns section, PAGE field', 'rowspan="2"' in allh and 'class="mc"' in allh and 'data-field="PAGE"' in allh)
    p = a.export('docx', 'preserve', 'T_roundtrip.docx'); d = pydocx.Document(p)
    check('DOCX round-trip: 2-column section kept', 'w:num="2"' in d.element.xml)
    for bad in ['corrupt.docx', 'old_format.docx', 'corrupt_truncated.pdf']:
        n0 = a.page.evaluate('APP.files.length'); msg = a.open(bad)
        check(f'error handling: {bad} rejected with message', a.page.evaluate('APP.files.length') == n0 and ('تالف' in msg or 'قديم' in msg), msg)

    # ---------- mixed / blank ----------
    a.open('mixed_text_and_scan.pdf'); a.run(); ps = a.pages()
    check('mixed PDF: page 1 hybrid text+OCR', ps[0]['meta']['method'] == 'hybrid' and 'Steel rebar' in ps[0]['text'] and '354,066.05' in ps[0]['text'])
    check('mixed PDF: scanned table rebuilt as ruled grid (6 cols)', a.page.evaluate("Math.max(...[...document.querySelectorAll('.pg[data-n=\"1\"] table')].map(t => VALID.grid(t).nc))") == 6)
    check('blank page detected', ps[1]['meta'].get('blank') is True)
    sp = a.export('spdf', fname='T_mixed_searchable.pdf')
    words = pdfplumber.open(sp).pages[0].extract_words()
    check('searchable PDF: OCR text layer over embedded scan', any(w['text'] == 'Steel' for w in words), len(words))

    # ---------- scanned, degraded ----------
    a.open('scanned_mixed_quality.pdf'); t0 = time.time(); a.run(); dt = time.time() - t0; ps = a.pages()
    check('scan: skew estimated (~1.6° page 1)', abs(abs(ps[0]['meta']['angle']) - 1.6) < 0.4, ps[0]['meta']['angle'])
    check('scan: upside-down page detected & corrected', ps[3]['meta']['rotK'] == 2)
    metrics = []
    for name, r, h in [('fin p1 200dpi skew', ref_fin[0], ps[0]['text']), ('invoice p1 skew', ref_inv[0], ps[1]['text']), ('fin p2 110dpi 2-col', ref_fin[1], ps[2]['text']), ('invoice p1 rotated 180', ref_inv[0], ps[3]['text'])]:
        rn, hn = collections.Counter(nums(r)), collections.Counter(nums(h)); tp = sum((rn & hn).values())
        m = {'page': name, 'CER': round(jiwer.cer(norm(r), norm(h)), 3), 'WER': round(jiwer.wer(norm(r), norm(h)), 3), 'num_recall': round(tp / max(1, sum(rn.values())), 3), 'num_precision': round(tp / max(1, sum(hn.values())), 3)}
        metrics.append(m); print('   ', m)
    check('scan: numeric recall ≥ 95% overall', sum(m['num_recall'] for m in metrics) / 4 >= 0.95, metrics)
    check('scan: 4 pages OCR under 120 s', dt < 120, round(dt, 1))

    # ---------- large ----------
    a.open('large_80pages.pdf'); t0 = time.time(); a.run(); dt = time.time() - t0
    check('80-page PDF processed < 60 s, all pages done', dt < 60 and all(p['status'] == 'done' for p in a.pages()), round(dt, 1))
    errs = [l for l in a.logs if 'pageerror' in l]
    check('no uncaught page errors', not errs, errs[:3])
finally:
    a.close()
json.dump({'results': R, 'ocr_metrics': metrics if 'metrics' in dir() else []}, open(os.path.join(OUT, 'results.json'), 'w'), ensure_ascii=False, indent=1)
print(f"\n{sum(r['ok'] for r in R)}/{len(R)} passed")
